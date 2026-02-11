import re
import math
import argparse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import os
import glob
from collections import defaultdict

# =========================
# MAPPING COMPLETO DELLE VARIAZIONI
# =========================

def get_task_order():
    """
    Ordine fisso dei task originali per la tabella Excel
    """
    return [
        "Open the middle layer of the drawer",
        "Put the bowl on the stove",
        "Put the wine bottle on the top of the drawer",
        "Open the top layer of the drawer and put the bowl inside",
        "Put the bowl on the top of the drawer",
        "Push the plate to the front of the stove",
        "Put the cream cheese in the bowl",
        "Turn on the stove",
        "Put the bowl on the plate",
        "Put the wine bottle on the rack"
    ]

def get_variation_mapping():
    """
    Mapping completo: ogni variazione → task originale
    Include: DEFAULT, L1, L2, L3
    Tutte le chiavi sono lowercase per matching case-insensitive
    """
    raw_mapping = {}
    
    # Task 1: Open the middle layer of the drawer
    orig_1 = "Open the middle layer of the drawer"
    raw_mapping["open the middle layer of the drawer"] = orig_1  # DEFAULT
    raw_mapping["Pull the middle layer of the drawer"] = orig_1  # L1
    raw_mapping["The middle layer of the drawer needs to be opened"] = orig_1  # L2
    raw_mapping["Open the layer of the drawer located between the top and bottom"] = orig_1  # L3
    
    # Task 2: Put the bowl on the stove
    orig_2 = "Put the bowl on the stove"
    raw_mapping["put the bowl on the stove"] = orig_2  # DEFAULT
    raw_mapping["Set the bowl on the stove"] = orig_2  # L1
    raw_mapping["The stove needs to have the bowl on it"] = orig_2  # L2
    raw_mapping["Put the object between the wine bottle and the cream cheese on the stove"] = orig_2  # L3
    
    # Task 3: Put the wine bottle on the top of the drawer
    orig_3 = "Put the wine bottle on the top of the drawer"
    raw_mapping["put the wine bottle on top of the drawer"] = orig_3  # DEFAULT
    raw_mapping["put the wine bottle on the top of the cabinet"] = orig_3  # VARIANTE
    raw_mapping["Place the wine bottle on the top of the drawer"] = orig_3  # L1
    raw_mapping["Top of the drawer needs to have the wine bottle on it"] = orig_3  # L2
    raw_mapping["Put the object behind the bowl on the top of the drawer"] = orig_3  # L3

    
    # Task 4: Open the top layer of the drawer and put the bowl inside
    orig_4 = "Open the top layer of the drawer and put the bowl inside"
    raw_mapping["Open the top drawer and put the bowl inside"] = orig_4  # DEFAULT
    raw_mapping["Open the top layer of the drawer and put the bowl inside"] = orig_4  # VARIANTE
    raw_mapping["Pull the top layer of the drawer and place the bowl inside"] = orig_4  # L1
    raw_mapping["Pull the top layer of the drawer and put the bowl inside"] = orig_4  # L1 alt
    raw_mapping["Store the bowl inside the top layer of the drawer"] = orig_4  # L2
    raw_mapping["Open the top layer of the drawer and put the object between the plate and the cream cheese inside"] = orig_4  # L3
    
    # Task 5: Put the bowl on the top of the drawer
    orig_5 = "Put the bowl on the top of the drawer"
    raw_mapping["Put the bowl on top of the drawer"] = orig_5  # DEFAULT
    raw_mapping["put the bowl on the top of the cabinet"] = orig_5  # VARIANTE
    raw_mapping["Place the bowl on the top of the drawer"] = orig_5  # L1
    raw_mapping["The top of the drawer needs to have the bowl on it"] = orig_5  # L2
    raw_mapping["Put the object between the wine bottle and the cream cheese on the top of the drawer"] = orig_5  # L3
    
    # Task 6: Push the plate to the front of the stove
    orig_6 = "Push the plate to the front of the stove"
    raw_mapping["push the plate to the front of the stove"] = orig_6  # DEFAULT
    raw_mapping["Move the plate to the front of the stove"] = orig_6  # L1
    raw_mapping["The space in front of the stove needs to have the plate in it"] = orig_6  # L2
    raw_mapping["Push the object in front of the drawer to the front of the stove"] = orig_6  # L3
    
    # Task 7: Put the cream cheese on the bowl
    orig_7 = "Put the cream cheese in the bowl"
    raw_mapping["Put the cream cheese in the bowl"] = orig_7  # DEFAULT
    raw_mapping["put the cream cheese on the bowl"] = orig_7  # VARIANTE
    raw_mapping["Place the cream cheese on the bowl"] = orig_7  # L1
    raw_mapping["Place the cream cheese in the bowl"] = orig_7  # L1 alt
    raw_mapping["The bowl needs to be filled with the cream cheese"] = orig_7  # L2
    raw_mapping["Put the object in front of the stove on the bowl"] = orig_7  # L3
    
    # Task 8: Turn on the stove
    orig_8 = "Turn on the stove"
    raw_mapping["turn on the stove"] = orig_8  # DEFAULT
    raw_mapping["Switch on the stove"] = orig_8  # L1
    raw_mapping["The stove needs to be turned on"] = orig_8  # L2
    raw_mapping["Turn on the object behind the cream cheese"] = orig_8  # L3
    
    # Task 9: Put the bowl on the plate
    orig_9 = "Put the bowl on the plate"
    raw_mapping["put the bowl on the plate"] = orig_9  # DEFAULT
    raw_mapping["Place the bowl on the plate"] = orig_9  # L1
    raw_mapping["The plate needs to have the bowl on it"] = orig_9  # L2
    raw_mapping["Put the object between the wine bottle and the cream cheese on the plate"] = orig_9  # L3
    
    # Task 10: Put the wine bottle on the rack
    orig_10 = "Put the wine bottle on the rack"
    raw_mapping["put the wine bottle on the rack"] = orig_10  # DEFAULT
    raw_mapping["Place the wine bottle on the rack"] = orig_10  # L1
    raw_mapping["The rack needs to be filled with the wine bottle in it"] = orig_10  # L2
    raw_mapping["Put the object behind the bowl on the rack"] = orig_10  # L3
    
    # Converti tutte le chiavi in lowercase per matching case-insensitive
    mapping = {k.lower(): v for k, v in raw_mapping.items()}
    
    return mapping

# =========================
# PARSING TXT FILES
# =========================

def parse_txt_file(filepath):
    """
    Parsa il file txt TinyVLA e estrae task, success rate ed episodi
    Formato: 
    Original Command: Put the wine bottle on the rack
    ...
    # episodes: 50
    # successes: 42 (84.0%)
    Task success rate: 0.8400
    """
    rates = {}
    episodes = {}
    task_order = []
    
    with open(filepath, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    current_task = None
    current_episodes = None
    
    for i, line in enumerate(lines):
        line = line.strip()
        
        # Cerca "Original Command:" o "Variation Command:"
        if "Original Command:" in line:
            current_task = line.split("Original Command:")[-1].strip()
        elif "Variation Command:" in line:
            current_task = line.split("Variation Command:")[-1].strip()
        
        # Cerca numero di episodi
        if current_task and "# episodes:" in line:
            match = re.search(r'# episodes:\s*(\d+)', line)
            if match:
                current_episodes = int(match.group(1))
        
        # Cerca il task success rate
        if current_task and "Task success rate:" in line:
            match = re.search(r'Task success rate:\s*([\d.]+)', line)
            if match:
                rate_value = float(match.group(1)) * 100  # Converti in percentuale
                
                if current_task not in rates:  # Evita duplicati
                    rates[current_task] = rate_value
                    episodes[current_task] = current_episodes if current_episodes else 50
                    task_order.append(current_task)
                
                # Reset
                current_task = None
                current_episodes = None
    
    return rates, episodes, task_order

def merge_txt_files(txt_files_list):
    """
    Combina più file .txt dello stesso seed
    Se un task appare più volte, calcola la media dei success rate
    """
    from collections import defaultdict
    
    task_rates_list = defaultdict(list)
    task_episodes = {}
    all_tasks = []
    
    for filepath in txt_files_list:
        rates, episodes, task_order = parse_txt_file(filepath)
        
        for task in task_order:
            if task not in task_rates_list:
                all_tasks.append(task)
            task_rates_list[task].append(rates[task])
            task_episodes[task] = 50  # Sempre 50 episodi
    
    # Calcola la media dei rate per task duplicati
    merged_rates = {}
    for task, rates_list in task_rates_list.items():
        merged_rates[task] = sum(rates_list) / len(rates_list)
    
    return merged_rates, task_episodes, all_tasks

# =========================
# EXCEL GENERATION
# =========================

def write_excel_comparison(output_xlsx, txt_files_by_seed):
    """
    Genera l'Excel con il formato richiesto
    txt_files_by_seed: dict con chiave seed (0,1,2) e valore lista di file paths
    """
    # Parsa i 3 seed (ognuno può avere più file)
    all_rates = []
    all_episodes = []
    task_order = None
    
    print("\n[INFO] Parsing e merge dei file txt...")
    for seed_idx in range(3):
        filepaths = txt_files_by_seed[seed_idx]
        print(f"  - Seed {seed_idx}: {len(filepaths)} file")
        for fp in filepaths:
            print(f"    • {os.path.basename(fp)}")
        
        rates, episodes, order = merge_txt_files(filepaths)
        all_rates.append(rates)
        all_episodes.append(episodes)
        if task_order is None:
            task_order = order
        print(f"    ✓ Trovati {len(rates)} task totali")
    
    # Ottieni il mapping variazione → originale
    variation_to_original = get_variation_mapping()
    
    # Crea mapping inverso: originale → variazione (trovata nei file)
    original_to_variation = {}
    print("\n[INFO] Identificazione variazioni nei file...")
    for task in task_order:
        task_lower = task.lower()
        if task_lower in variation_to_original:
            orig = variation_to_original[task_lower]
            original_to_variation[orig] = task
            print(f"  ✓ '{task}' → '{orig}'")
        else:
            print(f"  ⚠ Non mappata: '{task}'")
            original_to_variation[task] = task
    
    # Usa l'ordine fisso dei task
    fixed_task_order = get_task_order()
    
    # Crea workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"
    
    # Header
    ws.append([
        "Original Task Command",
        "Variation Task Command",
        "Success Rate (%) - Seed 0",
        "Task Completion - Seed 0",
        "Success Rate (%) - Seed 1",
        "Task Completion - Seed 1",
        "Success Rate (%) - Seed 2",
        "Task Completion - Seed 2",
        "Mean Success Rate (%) ± Std",
        "Mean Task Completion"
    ])
    
    # Formatta header
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
    
    # Statistiche per la riga finale
    all_seed_rates = [[], [], []]
    all_seed_completions = [[], [], []]
    
    print("\n[INFO] Generazione tabella Excel...")
    
    # Processa ogni task NELL'ORDINE FISSO
    for orig_task in fixed_task_order:
        # Trova la variazione corrispondente trovata nei file
        variation = original_to_variation.get(orig_task, orig_task)
        
        seed_rates = []
        seed_completions = []
        seed_success_counts = []
        seed_episode_counts = []
        
        # Estrai dati per ogni seed
        for seed_idx in range(3):
            if variation in all_rates[seed_idx]:
                rate = all_rates[seed_idx][variation]
                ep = all_episodes[seed_idx][variation]
                
                success_count = int(round(rate / 100 * ep))
                
                seed_rates.append(rate)
                seed_completions.append(f"{success_count}/{ep}")
                seed_success_counts.append(success_count)
                seed_episode_counts.append(ep)
                
                all_seed_rates[seed_idx].append(rate)
                all_seed_completions[seed_idx].append((success_count, ep))
            else:
                seed_rates.append(float('nan'))
                seed_completions.append("0/50")
                seed_success_counts.append(0)
                seed_episode_counts.append(50)
        
        # Calcola media E deviazione standard per il task
        valid_rates = [r for r in seed_rates if r == r]  # Escludi NaN
        if valid_rates:
            mean_rate = sum(valid_rates) / len(valid_rates)
            if len(valid_rates) > 1:
                std_rate = math.sqrt(sum((r - mean_rate) ** 2 for r in valid_rates) / (len(valid_rates) - 1))
            else:
                std_rate = 0.0
            mean_std_display = f"{mean_rate:.1f}% ± {std_rate:.1f}%"
        else:
            mean_rate = float('nan')
            std_rate = float('nan')
            mean_std_display = "nan%"
        
        # Mostra la variazione (trunca se troppo lunga)
        var_display = variation if len(variation) <= 50 else variation[:47] + "..."
        
        total_success = sum(seed_success_counts)
        total_episodes = sum(seed_episode_counts)

        # Calcola la media: successi medi su episodi medi
        avg_success = int(round(mean_rate / 100 * 50))  # Media basata sul mean_rate
        avg_episodes = 50  # Ogni seed ha sempre 50 episodi

        # Aggiungi riga
        ws.append([
            orig_task,
            var_display,
            f"{seed_rates[0]:.1f}%" if seed_rates[0] == seed_rates[0] else "nan%",
            seed_completions[0],
            f"{seed_rates[1]:.1f}%" if seed_rates[1] == seed_rates[1] else "nan%",
            seed_completions[1],
            f"{seed_rates[2]:.1f}%" if seed_rates[2] == seed_rates[2] else "nan%",
            seed_completions[2],
            mean_std_display,
            f"{avg_success}/{avg_episodes}"  # <-- ORA MOSTRA SU 50
        ])
            
    # ================= RIGA FINALE: Mean% ± Std% =================
    final_row = ["Mean% ± Std%", ""]
    
    # Per ogni seed
    for seed_idx in range(3):
        rates = all_seed_rates[seed_idx]
        
        if rates:
            mean = sum(rates) / len(rates)
            if len(rates) > 1:
                std = math.sqrt(sum((r - mean) ** 2 for r in rates) / (len(rates) - 1))
            else:
                std = 0.0
            
            total_success = sum(c[0] for c in all_seed_completions[seed_idx])
            total_episodes = sum(c[1] for c in all_seed_completions[seed_idx])
            
            final_row.append(f"{mean:.2f}% ± {std:.2f}%")
            final_row.append(f"{total_success}/{total_episodes}")
        else:
            final_row.append("N/A")
            final_row.append("0/0")
    
    # Media globale (standard VLA: std tra le medie dei seed, non tra tutti i valori)
    seed_means = []
    for seed_idx in range(3):
        if all_seed_rates[seed_idx]:
            seed_mean = sum(all_seed_rates[seed_idx]) / len(all_seed_rates[seed_idx])
            seed_means.append(seed_mean)
    
    if seed_means:
        global_mean = sum(seed_means) / len(seed_means)
        if len(seed_means) > 1:
            global_std = math.sqrt(
                sum((m - global_mean) ** 2 for m in seed_means) / (len(seed_means) - 1)
            )
        else:
            global_std = 0.0
        
        global_success = sum(sum(c[0] for c in all_seed_completions[i]) for i in range(3))
        global_episodes = sum(sum(c[1] for c in all_seed_completions[i]) for i in range(3))
        
        final_row.append(f"{global_mean:.2f}% ± {global_std:.2f}%")
        final_row.append(f"{global_success}/{global_episodes}")
    else:
        final_row.append("N/A")
        final_row.append("0/0")
    
    ws.append(final_row)
    
    # Formatta ultima riga
    last_row = ws.max_row
    for cell in ws[last_row]:
        cell.font = Font(bold=True)
    
    # Auto-adjust colonne
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column_letter].width = adjusted_width
    
    wb.save(output_xlsx)
    print(f"\n[OK] Excel salvato in: {output_xlsx}\n")

# =========================
# AUTO-DETECT FILES
# =========================

def find_txt_files_by_seed(base_dir, pattern_prefix="EVAL-libero_goal-tiny_vla"):
    """
    Trova automaticamente i file .txt per ogni seed
    Pattern: {pattern_prefix}_*_seed{0,1,2}_*.txt
    """
    txt_files_by_seed = {0: [], 1: [], 2: []}
    
    print(f"\n[INFO] Cercando file .txt in: {base_dir}")
    print(f"Pattern: {pattern_prefix}*seed*{{0,1,2}}*.txt")
    
    for seed_idx in range(3):
        # Cerca tutti i file per questo seed (supporta sia _ che - come separatore)
        pattern = os.path.join(base_dir, f"{pattern_prefix}*seed{seed_idx}*.txt")
        matches = glob.glob(pattern)
        
        # Ordina per nome file (task groups)
        matches.sort()
        
        txt_files_by_seed[seed_idx] = matches
        
        if matches:
            print(f"  ✓ Seed {seed_idx}: {len(matches)} file")
            for m in matches:
                print(f"    • {os.path.basename(m)}")
        else:
            print(f"  ✗ Seed {seed_idx}: nessun file trovato!")
            return None
    
    return txt_files_by_seed

# =========================
# MAIN
# =========================

def main():
    parser = argparse.ArgumentParser(
        description="Genera tabella Excel da file txt TinyVLA (supporta file multipli per seed)"
    )
    
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument(
        "--txt_dir",
        help="Directory contenente i file .txt (auto-detect per seed)"
    )
    group.add_argument(
        "--manual",
        action="store_true",
        help="Modalità manuale: specifica file per ogni seed"
    )
    
    parser.add_argument(
        "--pattern",
        default="EVAL-libero_goal-tiny_vla",
        help="Prefix pattern per auto-detect (default: EVAL-libero_goal-tiny_vla)"
    )
    parser.add_argument(
        "--seed0",
        nargs="+",
        help="File per seed 0 (modalità manuale)"
    )
    parser.add_argument(
        "--seed1",
        nargs="+",
        help="File per seed 1 (modalità manuale)"
    )
    parser.add_argument(
        "--seed2",
        nargs="+",
        help="File per seed 2 (modalità manuale)"
    )
    parser.add_argument(
        "output_xlsx",
        help="File Excel di output"
    )
    
    args = parser.parse_args()
    
    print("\n" + "="*60)
    print("  TINYVLA LIBERO TASK COMPARISON FROM TXT FILES")
    print("  (Supporting Multiple Files Per Seed)")
    print("="*60)
    
    # Determina i file da usare
    if args.manual:
        if not (args.seed0 and args.seed1 and args.seed2):
            print("\n[ERROR] Modalità manuale richiede --seed0, --seed1, --seed2!")
            exit(1)
        txt_files_by_seed = {
            0: args.seed0,
            1: args.seed1,
            2: args.seed2
        }
    else:
        txt_files_by_seed = find_txt_files_by_seed(args.txt_dir, args.pattern)
        if txt_files_by_seed is None:
            print("\n[ERROR] Non ho trovato file per tutti e 3 i seed!")
            exit(1)
    
    write_excel_comparison(args.output_xlsx, txt_files_by_seed)
    
    print("="*60)
    print("  COMPLETATO!")
    print("="*60 + "\n")

if __name__ == "__main__":
    main()